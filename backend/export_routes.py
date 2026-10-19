"""Authenticated Excel exports for the learning-center PostgreSQL database."""
from io import BytesIO
from datetime import date, datetime
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from auth import require_auth
from db import get_connection

router = APIRouter(prefix="/export", dependencies=[Depends(require_auth)])
PACIFIC = ZoneInfo("America/Los_Angeles")


def fetch_rows(cur, query, params=()):
    cur.execute(query, params)
    columns = [column.name for column in cur.description]
    return columns, cur.fetchall()


def safe_cell(value):
    if isinstance(value, datetime):
        # Excel does not accept timezone-aware datetime objects.
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r", "\n")):
        # Prevent Excel from evaluating untrusted student text as a formula.
        return "'" + value
    return value


def workbook_response(sheets, filename):
    wb = Workbook(write_only=False)
    wb.remove(wb.active)
    for title, (columns, rows) in sheets:
        ws = wb.create_sheet(title)
        ws.append(columns)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for row in rows:
            ws.append([safe_cell(value) for value in row])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for index, name in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(index)].width = min(max(len(name) + 3, 16), 38)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"', "Cache-Control": "no-store"},
    )


def fetch_sessions_with_names(cur, where_clause="", params=()):
    """Include names first, preserve all original session columns afterward."""
    # Use database metadata so adding new session fields won't silently remove
    # them from future exports. Move the identifying IDs to the end.
    cur.execute("SELECT * FROM sessions LIMIT 0")
    original_columns = [column.name for column in cur.description]
    detail_columns = [name for name in original_columns if name not in ("session_id", "student_id")]
    id_columns = [name for name in ("session_id", "student_id") if name in original_columns]
    # Column names are obtained from PostgreSQL schema metadata, not user input.
    from psycopg import sql
    selected = [
        sql.SQL('s.first_name AS "First Name"'),
        sql.SQL('s.last_name AS "Last Name"'),
    ] + [sql.SQL('sess.{}').format(sql.Identifier(name)) for name in detail_columns + id_columns]
    query = sql.SQL("SELECT {} FROM sessions AS sess JOIN students AS s ON s.student_id = sess.student_id {} ORDER BY sess.session_date DESC, sess.session_id DESC").format(
        sql.SQL(", ").join(selected), sql.SQL(where_clause)
    )
    return fetch_rows(cur, query, params)


@router.get("/today")
def export_today():
    today = datetime.now(PACIFIC).date()
    with get_connection() as conn:
        with conn.cursor() as cur:
            session_columns, sessions = fetch_sessions_with_names(
                cur,
                "WHERE (sess.session_date AT TIME ZONE 'America/Los_Angeles')::date = %s",
                (today,),
            )
            student_columns, students = fetch_rows(
                cur,
                """SELECT * FROM students
                   WHERE student_id IN (
                       SELECT DISTINCT student_id FROM sessions
                       WHERE (session_date AT TIME ZONE 'America/Los_Angeles')::date = %s
                   )
                   ORDER BY last_name, first_name, student_id""",
                (today,),
            )
    return workbook_response(
        [("Today's Students", (student_columns, students)), ("Today's Sessions", (session_columns, sessions))],
        f"student_attendance_{today.isoformat()}.xlsx",
    )


@router.get("/all")
def export_all():
    with get_connection() as conn:
        with conn.cursor() as cur:
            student_columns, students = fetch_rows(
                cur, "SELECT * FROM students ORDER BY last_name, first_name, student_id"
            )
            session_columns, sessions = fetch_sessions_with_names(cur)
    today = datetime.now(PACIFIC).date()
    return workbook_response(
        [("Students", (student_columns, students)), ("Sessions", (session_columns, sessions))],
        f"student_progress_full_export_{today.isoformat()}.xlsx",
    )
